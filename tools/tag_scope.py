"""Add a tag to the Scope column (J) of test-plan rows: "Switch + Compute" -> "Switch + Compute; IFCP". The run form
and tools/select_cases.py match one tag or the whole cell, so --scope IFCP picks every tagged case and the old
scope values keep working. Idempotent (a row that already has the tag is left), --remove takes it off again.
    python tools/tag_scope.py testplan/full_rack_testplan.xlsx --tag IFCP --cases testplan/tags/ifcp.txt [--dry-run]
The workbook must be closed in Excel (refused while its ~$ lock file exists); a timestamped backup is written next to it.
"""
from __future__ import annotations

import argparse
import shutil
import sys
import time
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from select_cases import TAG_SEP, check_layout, tags  # noqa: E402
from update_testplan import find_row  # noqa: E402

SCOPE_COL = 10


def read_cases(path: Path) -> list[str]:
    out = []
    for line in path.read_text().splitlines():
        line = line.split("#", 1)[0]
        out += [t for t in line.replace(",", " ").split() if t]
    return list(dict.fromkeys(out))


def retag(cell, tag: str, remove: bool) -> str:
    have = tags(cell)
    if remove:
        new = [t for t in have if t.lower() != tag.lower()]
    else:
        new = have if any(t.lower() == tag.lower() for t in have) else have + [tag]
    return f"{TAG_SEP} ".join(new)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workbook", type=Path)
    ap.add_argument("--tag", required=True)
    ap.add_argument("--cases", type=Path, required=True, help="file of plan IDs (space / newline separated, # comments)")
    ap.add_argument("--remove", action="store_true", help="take the tag off instead")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--no-backup", action="store_true")
    a = ap.parse_args(argv)
    if TAG_SEP in a.tag or not a.tag.strip():
        ap.error(f"a tag cannot be empty or contain {TAG_SEP!r}")
    lock = a.workbook.with_name("~$" + a.workbook.name)
    if lock.exists() and not a.dry_run:
        print(f"refusing: {a.workbook.name} is open in Excel ({lock.name} exists) - save and close it first")
        return 2
    wb = openpyxl.load_workbook(a.workbook)
    ws = wb["Testcases"]
    check_layout(ws)
    changed, same, missing = [], 0, []
    for tc in read_cases(a.cases):
        r = find_row(ws, tc)
        if r is None:
            missing.append(tc)
            continue
        old = ws.cell(r, SCOPE_COL).value
        new = retag(old, a.tag, a.remove)
        if (old or "") == new:
            same += 1
            continue
        changed.append((tc, old, new))
        ws.cell(r, SCOPE_COL).value = new
    for tc, old, new in changed:
        print(f"  {tc:<12} {old or '(empty)'!s:<36} -> {new}")
    print(f"{len(changed)} changed, {same} already {'without' if a.remove else 'with'} the tag"
          + (f", not in the workbook: {' '.join(missing)}" if missing else ""))
    if a.dry_run or not changed:
        return 1 if missing else 0
    if not a.no_backup:
        bak = a.workbook.with_name(f"{a.workbook.stem}.backup-{time.strftime('%Y%m%d-%H%M%S')}{a.workbook.suffix}")
        shutil.copy2(a.workbook, bak)
        print(f"backup: {bak}")
    wb.save(a.workbook)
    print(f"saved {a.workbook}")
    return 1 if missing else 0


if __name__ == "__main__":
    sys.exit(main())
