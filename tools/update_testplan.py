"""Write helios-tests run results into the test-plan workbook (Testcases sheet).

    python tools/update_testplan.py testplan/full_rack_testplan.xlsx results/runs/<rack>/<run> [more runs] \\
        --sw-version "AFM 0.1.0-AICC-817" --label "g02-g17 full rack, 2026-10-14" [--dry-run]

Per plan ID (several runs merge like tools/html_report.py: a run that executed a test beats one that
skipped it, otherwise the later run wins):
  FAIL     any item failed or errored       -> FAIL, notes = first failure reasons
           ... but every failure is the test ENVIRONMENT (tool not installed, permission denied, unreachable,
           name lookup, GPU not openable)    -> Pending, notes "Environment: <reason>" (not a product verdict)
  PARTIAL  passed only on the trays the pre-flight kept -> PARTIAL, notes list the excluded trays
  PASS     every executed item passed       -> PASS (notes list items that were skipped, if any)
  BLOCKED  a pre-flight refused to start    -> Pending, notes "Blocked: <reason>"
  SKIP     nothing executed                 -> Skipped, notes "Skipped: <reason>" (--no-skipped leaves the row)
New notes go ABOVE the row's earlier notes (kept as "Earlier: ..."); a rerun with the same --label replaces its
own previous note. Columns F (status), G (SW version), H (notes) and an empty Scope column are written; the workbook must be
closed in Excel (refused while its ~$ lock file exists). A timestamped backup is written next to it
(--no-backup to skip). Execution Summary formulas recalculate when Excel opens the file.
"""
from __future__ import annotations

import argparse
import copy
import shutil
import sys
import time
from collections import OrderedDict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment

sys.path.insert(0, str(Path(__file__).resolve().parent))
from html_report import load_runs  # noqa: E402

# test IDs whose workbook row has (or, before 21 Oct 2026, had) a shared / missing TID: test ID -> (TID in the
# BKC-26.12 plan, title fragment). TC-AFM-005A / 048A now have their own rows; the entries keep older copies working.
ROW_ALIASES = {
    "TC-AFM-005A": ("TC-AFM-005", "RCCL Cross Domain"),
    "TC-AFM-005": ("TC-AFM-005", "Failover Policy CRUD"),
    "TC-AFM-048A": ("TC-AFM-048", "Audit Events"),
    "TC-AFM-048": ("TC-AFM-048", "Techsupport"),
    "TC-AFM-035A": (None, "Not allowed in Minirack"),
}
COL_STATUS, COL_VERSION, COL_NOTES, COL_ACCESS, COL_SCOPE = 6, 7, 8, 9, 10

# failures that say "the test could not run here", not "the product is broken"
ENV_PATTERNS = ("command not found", "not installed", "permission denied", "no rocm-capable device",
                "name or service not known", "unreachable", "connect failed", "unable to open channel",
                "no such file or directory", "rankaborted", "killed after the rank on")


def reason(r: dict) -> str:
    m = (r.get("message") or "").strip()
    if r["outcome"] == "skipped":
        return m.removeprefix("Skipped: ")[:250]
    lines = [ln.strip(" -") for ln in m.splitlines() if ln.strip()]
    detail = [ln for ln in lines if "check(s) failed" not in ln and "pre-flight failed" not in ln]
    return ("; ".join(detail[:3]) or (lines[0] if lines else r["outcome"]))[:400]


def verdicts(items: list[dict]) -> "OrderedDict[str, tuple[str, str]]":
    by: "OrderedDict[str, list[dict]]" = OrderedDict()
    for r in items:
        if r.get("tc_id"):
            by.setdefault(r["tc_id"], []).append(r)
    out = OrderedDict()
    for tc, rows in by.items():
        outs = [r["outcome"] for r in rows]
        bad = [r for r in rows if r["outcome"] in ("failed", "error")]
        if bad and all(_env_only(r) for r in bad):
            out[tc] = ("ENV", "Environment: " + reason(bad[0]))
        elif bad:
            out[tc] = ("FAIL", reason(bad[0]))
        elif "blocked" in outs:
            out[tc] = ("BLOCKED", "Blocked: " + reason(next(r for r in rows if r["outcome"] == "blocked")))
        elif "partial" in outs:
            out[tc] = ("PARTIAL", reason(next(r for r in rows if r["outcome"] == "partial")))
        elif "passed" in outs:
            skipped = [r["test"] for r in rows if r["outcome"] == "skipped"]
            out[tc] = ("PASS", f"{len(rows) - len(skipped)} item(s) passed"
                       + (f"; skipped: {', '.join(skipped)[:200]}" if skipped else ""))
        else:
            out[tc] = ("SKIP", "Skipped: " + reason(rows[0]))
    return out


def _env_only(r: dict) -> bool:
    """Every failure line of this item is an environment problem (see ENV_PATTERNS)."""
    lines = [ln.strip(" -").lower() for ln in (r.get("message") or "").splitlines()
             if ln.strip() and "check(s) failed" not in ln]
    return bool(lines) and all(any(p in ln for p in ENV_PATTERNS) for ln in lines)


def merge_notes(new: str, old, label: str) -> str:
    """New note first, earlier notes kept below; a previous note from the same label is replaced."""
    old = str(old or "").strip()
    if label:   # drop every earlier note from this same label, wherever it sits in the history
        paras = [x.removeprefix("Earlier: ") for x in old.split("\n\n")]
        old = "\n\n".join(x for x in paras if x and not x.startswith(f"[{label}]"))
    return f"{new}\n\nEarlier: {old}" if old else new


def find_row(ws, tc: str) -> int | None:
    # the case's own TID first (TC-AFM-005A / 048A have their own rows since 21 Oct 2026); else the shared BKC-26.12
    # TID + title fragment of ROW_ALIASES, for workbook copies made before the rename
    for r in range(3, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if (a.strip() if isinstance(a, str) else a) == tc:
            frag = ROW_ALIASES.get(tc, (None, None))[1]      # a case on a once-shared TID must match its title too
            if frag is None or frag.lower() in str(ws.cell(r, 3).value or "").lower():
                return r
    tid, frag = ROW_ALIASES.get(tc, (tc, None))
    for r in range(3, ws.max_row + 1):
        a = ws.cell(r, 1).value
        a = a.strip() if isinstance(a, str) else a
        if a != tid:
            continue
        if frag is None or frag.lower() in str(ws.cell(r, 3).value or "").lower():
            return r
    return None


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workbook", type=Path)
    ap.add_argument("runs", type=Path, nargs="+", help="results/runs/<rack>-<timestamp> folders")
    ap.add_argument("--sw-version", default="", help="SW Version column text (e.g. 'AFM 0.1.0-AICC-817')")
    ap.add_argument("--label", default="", help="prefix for notes, e.g. 'g02-g17 full rack, 2026-10-14'")
    ap.add_argument("--mark-skipped", action="store_true", help=argparse.SUPPRESS)   # now the default
    ap.add_argument("--no-skipped", action="store_true", help="leave rows of not-executed IDs untouched")
    ap.add_argument("--dry-run", action="store_true", help="print what would change, write nothing")
    ap.add_argument("--no-backup", action="store_true")
    a = ap.parse_args()

    wb_path = a.workbook
    lock = wb_path.with_name("~$" + wb_path.name)
    if lock.exists() and not a.dry_run:
        print(f"refusing: {wb_path.name} is open in Excel ({lock.name} exists) - save and close it first")
        return 2
    items = load_runs(a.runs)
    res = verdicts(items)
    wb = openpyxl.load_workbook(wb_path)
    ws = wb["Testcases"]
    ref = next((r for r in range(3, ws.max_row + 1) if ws.cell(r, COL_STATUS).value), 3)  # style donor
    changed, unmatched, counts = [], [], {}
    runs = ", ".join(sorted({r["_run"] for r in items}))
    for tc, (v, note) in res.items():
        if v == "SKIP" and a.no_skipped:
            counts["SKIP (untouched)"] = counts.get("SKIP (untouched)", 0) + 1
            continue
        row = find_row(ws, tc)
        if row is None:
            unmatched.append(tc)
            continue
        status = {"BLOCKED": "Pending", "ENV": "Pending", "SKIP": "Skipped"}.get(v, v)
        text = " ".join(x for x in (f"[{a.label}]" if a.label else "", note, f"(helios-tests {runs})") if x)
        old = ws.cell(row, COL_STATUS).value
        changed.append((tc, row, old, status, text))
        counts[v] = counts.get(v, 0) + 1
        if not a.dry_run:
            ws.cell(row, COL_STATUS).value = status
            if a.sw_version:
                ws.cell(row, COL_VERSION).value = a.sw_version
            ws.cell(row, COL_NOTES).value = merge_notes(text, ws.cell(row, COL_NOTES).value,
                                                        a.label)
            if str(ws.cell(2, COL_SCOPE).value or "").strip() == "Scope" and not ws.cell(row, COL_SCOPE).value:
                from access_needed import scope
                ws.cell(row, COL_SCOPE).value = scope(str(ws.cell(row, COL_ACCESS).value or ""))
            for c in (COL_STATUS, COL_VERSION, COL_NOTES):
                ws.cell(row, c)._style = copy.copy(ws.cell(ref, c)._style)
            ws.cell(row, COL_NOTES).alignment = Alignment(wrap_text=True, vertical="top")
    for tc, row, old, new, text in changed:
        print(f"{tc:<18} row {row:<4} {str(old or '-'):<8} -> {new:<8} {text[:110]}")
    print(f"\n{len(changed)} row(s) {'would change' if a.dry_run else 'updated'}: {counts}")
    if unmatched:
        print(f"no workbook row for: {sorted(unmatched)}")
    if a.dry_run:
        return 0
    if not a.no_backup:
        bak = wb_path.with_name(f"{wb_path.stem}.backup-{time.strftime('%Y%m%d-%H%M%S')}{wb_path.suffix}")
        shutil.copy2(wb_path, bak)
        print(f"backup: {bak}")
    wb.calculation.fullCalcOnLoad = True
    wb.save(wb_path)
    print(f"saved {wb_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
