"""Add / refresh an "Access needed" column in the test-plan workbook (Testcases column I, mortaro column H):
what a run of each case needs - AFM API, switch SSH, compute SSH, datapath traffic, BMC/PDU, a special rack
or server, or manual work.

    python tools/access_needed.py testplan/full_rack_testplan.xlsx [more workbooks] [--dry-run]
    python tools/access_needed.py --summary          # counts per access class, no workbook needed

Automated cases are classified from the code (which suite, disruptive or not); the rest from the plan
category and title (rules below, marked "(est.)" when inferred from the title only). The workbook must be
closed in Excel (refused while its ~$ lock file exists); a timestamped backup is written next to it.
"""
from __future__ import annotations

import argparse
import ast
import copy
import re
import shutil
import sys
import time
from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path(__file__).resolve().parent))
from plan_cases import load_cases  # noqa: E402
from update_testplan import ROW_ALIASES  # noqa: E402

AFM = "AFM API"
AFM_SW = "AFM API + switch SSH"
SW = "Switch SSH"
SW_CMP = "Switch + compute SSH"
CMP = "Compute SSH"
TRAFFIC = "Compute SSH + datapath traffic"

# automated suites -> access (by test file); per-ID overrides below
FILE_ACCESS = {
    "tests/l1/test_switch_hw.py": SW,
    "tests/switch/test_switch_agent.py": SW,
    "tests/l1/test_afm_datapath.py": TRAFFIC,
    "tests/l1/test_rccl_smoke.py": TRAFFIC,
    "tests/ifoe/test_ifoe_driver.py": TRAFFIC,
    "tests/ras/test_ras_reset.py": TRAFFIC,
    "tests/telemetry/test_metrics_exporter.py": CMP,
    "tests/platform/test_node_platform.py": SW_CMP,
    "tests/sideband/test_sideband_bringup.py": SW_CMP,
    "tests/sideband/test_ifcp.py": SW_CMP,
    "tests/scaleout/test_rdma_perftest.py": "Compute SSH + scale-out NICs",
}
ID_ACCESS = {
    "TC-HW-002": SW, "TC-HW-008": TRAFFIC, "TC-HW-013": SW_CMP,
    "TC-RAS-reset-007": TRAFFIC + " + node reboot", "TC-RAS-reset-008": TRAFFIC + " + BMC",
    "TC-RAS-reset-009": TRAFFIC + " + PDU",
    "TC-CTL-059": SW, "TC-BUG-022": SW, "TC-E2E-118": SW, "TC-BUG-021": SW + " + switch reboot",
}
# automated cases that also need lab equipment (they skip until the inventory provides it)
EQUIPMENT = {
    "tests/afm/test_afm_mch.py": "MCH rack (18-link)",
    "TC-AFM-044": "LDAP server", "TC-CTL-119": "RADIUS server", "TC-CTL-042": "LDAP + RADIUS servers",
    "TC-CTL-101": "syslog server",
    **{t: "upgrade images" for t in ("TC-CTL-054", "TC-CTL-078", "TC-CTL-127", "TC-SWA-128", "TC-SWA-164",
                                     "TC-HW-022", "TC-AFM-UPGRADE-004")},
    **{t: "out-of-band mgmt network control" for t in ("TC-AFM-008", "TC-E2E-093", "TC-BUG-003", "TC-CTL-050",
                                                       "TC-BUG-019")},
}

# not automated: (category, title regex) -> access, first match wins
RULES = [
    ("mch", r".", "MCH rack (18-link) + AFM API + SSH"),
    ("controller", r"arranged|tooltip|loading time|Hide Menu|Time Zone|timezone", "Manual (AFM UI)"),
    ("controller", r"LDAP|RADIUS|Radius", "AFM API + LDAP/RADIUS server"),
    ("controller", r"Syslog", "AFM API + syslog server"),
    ("controller", r"certificate|Mutual TLS|SPDM|signed certificate", "Switch SSH + certificates/PKI"),
    ("controller", r"upgrade|firmware|cluster move", "Switch SSH + upgrade images"),
    ("controller", r"network disruption", "AFM API + switch SSH + out-of-band network control"),
    ("controller", r"Installation|bootstrap|selection criteria", "Switch SSH (fresh AFM install)"),
    ("controller", r"NMD", "Switch SSH (agent internals)"),
    ("controller", r"techsupport", AFM),
    ("controller", r"[Ll]ogs", SW),
    ("afm", r"Management Network", "Compute + switch SSH + mgmt network control"),
    ("afm", r"RADIUS|LDAP", "AFM API + LDAP/RADIUS server"),
    ("afm", r"Compute Link Flap|afmctl|station", "AFM API + compute SSH (afmctl)"),
    ("afm-upgrade", r".", "Switch SSH + upgrade image"),
    ("afm-reconfigure", r"Accelerator Reset", "AFM API + compute SSH (afmctl)"),
    ("e2e-failover", r"ping check|compute side|management connectivity|Forwarding table|config phase|node manager|services/agents",
     "AFM API + compute SSH (afmctl)"),
    ("e2e-failover", r"status of GPU/Switch ports", AFM_SW),
    ("e2e-failover", r".", AFM + " (destructive; procedure unclear)"),
    ("bug-regression", r"management IP interface", "AFM API + switch SSH + mgmt network control"),
    ("bug-regression", r"Fboss", "Fboss switch"),
    ("bug-regression", r"[Pp]assword recovery", "Switch SSH + password-recovery procedure"),
    ("bug-regression", r"fru\.json|eth1|non-default interface", SW),
    ("bug-regression", r"Lanebitmask|ready state|programming errors|sysfs|Nodeagent|incompatible", "AFM API + compute SSH (afmctl)"),
    ("bug-regression", r".", AFM_SW),
    ("event-hub", r".", "AFM API + switch SSH (controller internals)"),
    ("stress", r".", "AFM API + switch SSH (full rack, long run)"),
    ("ras-reset", r".", "AFM API + compute SSH + datapath traffic"),
    ("ifcp-version", r".", SW_CMP + " + packet manipulation"),
    ("aggwatch", r"per-object|Per-object", AFM),
    ("hardware", r".", SW_CMP + " (hardware access)"),
    ("scale-out-rdma", r".", "Compute SSH + scale-out NICs"),
    ("switch-hw", r".", SW),
    ("switch-agent", r".", SW),
    (None, r".", SW_CMP),   # sideband / IFCP / node-agent / agent-restart and anything else
]


# what a test (or the module helpers it calls) drives besides the AFM API
SSH_MARKERS = re.compile(r"\b(links|rack|ssh_switch|ssh_by_ip|slot_host|sb_agents|afm_agent|signal_sb_agent|restart_controller|"
                         r"controller_action|controller_state|controllers|switch_agent_pid|epoch|controller_logs_since|"
                         r"vlan_count|static_macs|gpu_links|link_candidates|port_admin|reset_switch_identity|_need_sideband|"
                         r"_vlans)\b")
COMPUTE_MARKERS = re.compile(r"\bops\.(tray|ifcp|ifcp_up|compute_port_admin|station_map|tray_time|reset_logged_since|"
                             r"phases|afmctl)\b|\b(ComputePorts|_mac_counters|_afm_names)\b|\btray\.run\b")
TRAFFIC_MARKERS = re.compile(r"\b(datapath|_traffic|_ras|_vpod_split)\b")


def _uses(fn: ast.FunctionDef, funcs: dict[str, ast.FunctionDef], rx, seen=None) -> bool:
    seen = seen if seen is not None else set()
    if fn.name in seen:
        return False
    seen.add(fn.name)
    body = [b for b in fn.body if not (isinstance(b, ast.Expr) and isinstance(getattr(b, "value", None), ast.Constant)
                                       and isinstance(b.value.value, str))]          # drop the docstring
    code = "\n".join(ast.unparse(b) for b in body)
    code = re.sub(r"(['\"]).*?\1", "''", code)                                        # and string literals
    if rx.search(code):
        return True
    called = {n.func.id for n in ast.walk(fn) if isinstance(n, ast.Call) and isinstance(n.func, ast.Name)}
    return any(_uses(funcs[c], funcs, rx, seen) for c in called if c in funcs)


def automated() -> dict[str, tuple[str, bool, str]]:
    """{tc id: (test file, disruptive?, extra access label)} from the test sources."""
    out = {}
    helper_src = {}
    for h in (ROOT / "tests").rglob("afm_helpers.py"):
        helper_src.update({n.name: n for n in ast.walk(ast.parse(h.read_text())) if isinstance(n, ast.FunctionDef)})
    for f in sorted((ROOT / "tests").rglob("test_*.py")):
        src = f.read_text()
        tree = ast.parse(src)
        funcs = dict(helper_src)
        funcs.update({n.name: n for n in tree.body if isinstance(n, ast.FunctionDef)})
        mod_disr = any("disruptive" in ln for ln in src.splitlines() if ln.startswith("pytestmark"))
        mod_traffic = any("datapath_ready" in ln for ln in src.splitlines() if ln.startswith("pytestmark"))
        for node in ast.walk(tree):
            if isinstance(node, ast.FunctionDef) and node.name.startswith("test_"):
                decos = [ast.unparse(d) for d in node.decorator_list]
                disr = mod_disr or any("disruptive" in d for d in decos)
                args = {a.arg for a in node.args.args}
                sw = bool(args & {"links", "rack"}) or _uses(node, funcs, SSH_MARKERS)
                cmp_ = bool(args & {"computes"}) or _uses(node, funcs, COMPUTE_MARKERS)
                traffic = mod_traffic or _uses(node, funcs, TRAFFIC_MARKERS)
                parts = [p for p, on in (("switch", sw), ("compute", cmp_ or traffic)) if on]
                extra = (" + " + " + ".join(parts) + " SSH") if parts else ""
                extra += " + datapath traffic" if traffic else ""
                if not args & {"afm", "afm_anon", "afm_cfg", "ops", "afm_ready", "make_user", "watch"}:
                    extra = "NOAFM" + extra          # no AFM call at all: SSH / traffic only
                for d in decos:
                    m = re.search(r"mark\.tc\(['\"]([^'\"]+)", d)
                    if m:
                        out.setdefault(m.group(1), (str(f.relative_to(ROOT)), disr, extra))
    return out


def classify_auto(f: str, disr: bool, extra: str) -> str:
    if f.startswith("tests/afm/"):
        if extra.startswith("NOAFM"):
            return extra[len("NOAFM + "):].replace("compute SSH", "Compute SSH").replace("switch", "Switch", 1)
        return AFM + extra if extra else (AFM + " (disruptive)" if disr else AFM)
    return FILE_ACCESS.get(f, SW_CMP)


def classify(case: dict, auto: dict) -> str:
    tid = case["id"]
    if tid in ID_ACCESS:
        return ID_ACCESS[tid]
    if tid in auto:
        f, disr, extra = auto[tid]
        need = EQUIPMENT.get(tid) or EQUIPMENT.get(f)
        if need:
            return f"{need} + " + classify_auto(f, disr, extra)
        return classify_auto(f, disr, extra)
    title = case.get("title") or ""
    for cat, rx, access in RULES:
        if (cat is None or case.get("category") == cat) and re.search(rx, title):
            return access + ("" if access.startswith(("MCH", "Manual")) or cat in ("stress", "event-hub") else " (est.)")
    return SW_CMP + " (est.)"


def scope(access: str) -> str | None:
    """Scope column (J): where a case runs - Switch, Compute, Switch + Compute, or Switch (AFM API) (the AFM
    controller runs on the switch trays) - derived from its Access needed value."""
    a = (access or "").lower()
    sw, cp = "switch" in a or a.startswith("mch"), "compute" in a or a.startswith("mch")
    if sw and cp:
        return "Switch + Compute"
    if sw or cp:
        return "Switch" if sw else "Compute"
    return "Switch (AFM API)" if "afm" in a else None


def table() -> dict[str, str]:
    cases = load_cases()
    auto = automated()
    return {c["id"]: classify(c, auto) for c in cases}


def rack_validated(run_dirs: list[Path]) -> set[str]:
    """Plan IDs that were run on a real rack in any of these runs - whatever the verdict: a case skipped on the rack
    for a platform reason (e.g. no PSU fans on a busbar-fed switch) is still automated and attempted."""
    import csv

    out = set()
    for d in run_dirs:
        for f in list(d.glob("tc_summary.csv")) + list(d.glob("*/tc_summary.csv")):
            for r in csv.DictReader(open(f)):
                if r["verdict"] in ("PASS", "FAIL", "PARTIAL", "SKIP", "BLOCKED"):
                    out.add(r["tc_id"])
    return out


def automation_values(validated: set[str]) -> dict[str, str]:
    """Automaton(Yes/No/Inprogress): Yes = automated and run on a rack (any result, including Skipped); Inprogress =
    automated, never run on a rack; No = not automated."""
    auto = automated()
    cases = load_cases()
    return {c["id"]: ("Yes" if c["id"] in validated else "Inprogress") if c["id"] in auto else "No" for c in cases}


def write(wb_path: Path, access: dict[str, str], dry: bool, automation: dict[str, str] | None = None) -> None:
    lock = wb_path.with_name("~$" + wb_path.name)
    if lock.exists() and not dry:
        print(f"refusing: {wb_path.name} is open in Excel ({lock.name} exists) - save and close it first")
        return
    wb = openpyxl.load_workbook(wb_path)
    alias_of = {}
    for test_id, (tid, frag) in ROW_ALIASES.items():
        alias_of[(tid, frag)] = test_id
    counts, missing = Counter(), []
    for sheet, col, hdr_ref in (("Testcases", 9, 8), ("mortaro", 8, 7)):
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        ws.cell(2, col).value = "Access needed"
        ws.cell(2, col)._style = copy.copy(ws.cell(2, hdr_ref)._style)
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 42
        for r in range(3, ws.max_row + 1):
            tid = ws.cell(r, 1).value
            tid = tid.strip() if isinstance(tid, str) else tid
            title = str(ws.cell(r, 3 if sheet == "Testcases" else 7).value or "")
            if not tid and sheet == "Testcases" and "Not allowed in Minirack" in title:
                tid = "TC-AFM-035A"
            if not tid or not str(tid).startswith("TC-"):
                continue
            key = next((t for (a, frag), t in alias_of.items() if a == tid and frag.lower() in title.lower()), tid)
            val = access.get(key) or access.get(tid)
            if val is None:
                missing.append(tid)
                continue
            counts[val.split(" (")[0]] += 1
            if automation is not None and sheet == "Testcases":
                av = automation.get(key) or automation.get(tid)
                if av:
                    counts[f"Automation={av}"] += 1
                    if not dry:
                        ws.cell(r, 4).value = av
            if not dry:
                ws.cell(r, col).value = val
                ws.cell(r, col)._style = copy.copy(ws.cell(r, hdr_ref)._style)
                ws.cell(r, col).alignment = Alignment(wrap_text=True, vertical="top")
    print(f"{wb_path.name}: {sum(v for k, v in counts.items() if not k.startswith('Automation='))} rows "
          f"{'would get' if dry else 'got'} an access value")
    for k, n in counts.most_common():
        print(f"  {n:>4}  {k}")
    if missing:
        print(f"  no plan entry for: {sorted(set(missing))[:10]}")
    if dry:
        return
    bak = wb_path.with_name(f"{wb_path.stem}.backup-{time.strftime('%Y%m%d-%H%M%S')}{wb_path.suffix}")
    shutil.copy2(wb_path, bak)
    wb.save(wb_path)
    print(f"  saved (backup {bak.name})")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workbooks", type=Path, nargs="*")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--summary", action="store_true", help="print counts per access class and exit")
    ap.add_argument("--automation", nargs="*", type=Path, metavar="RUN_DIR",
                    help="also rewrite column D (Automaton Yes/No/Inprogress) from the code; Yes = run on a rack (any "
                         "result, Skipped included) in one of these result folders (e.g. results/g02-g17-*)")
    a = ap.parse_args()
    access = table()
    if a.summary or not a.workbooks:
        auto = automated()
        for label, ids in (("automated", [i for i in access if i in auto]), ("not automated", [i for i in access if i not in auto])):
            print(f"{label} ({len(ids)}):")
            for k, n in Counter(access[i].replace(" (est.)", "") for i in ids).most_common():
                print(f"  {n:>4}  {k}")
        return 0
    automation = automation_values(rack_validated(a.automation)) if a.automation is not None else None
    for wb in a.workbooks:
        write(wb, access, a.dry_run, automation)
    return 0


if __name__ == "__main__":
    sys.exit(main())
