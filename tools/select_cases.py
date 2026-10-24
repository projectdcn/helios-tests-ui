"""Build a --tc list from the test-plan workbook's columns (Level, Scope, Disruptive, Origin, status) - e.g. every
non-disruptive automated case whose scope is the switch or the AFM API, or the health-check cases (--health
Blocker,Warning) - split into the AFM suite (needs the AFM password) and the SSH suites.

    python tools/select_cases.py testplan/full_rack_testplan.xlsx --disruptive Non-disruptive \\
        --scope "Switch,Switch (AFM API)" --suite afm -o results/notes/afm-nondisruptive-tcs.txt
    .venv/bin/pytest tests/afm --inventory <inv> --tc "$(cat results/notes/afm-nondisruptive-tcs.txt)"

Only automated cases are listed (a test with that plan ID exists). Filters are comma-separated and exact
(case-insensitive; a comma inside parentheses is part of the value); "(empty)" matches a blank cell, e.g.
--status "(empty)" = not run yet; --suite: afm (tests/afm), ssh (everything else), all.
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from access_needed import automated  # noqa: E402
from update_testplan import ROW_ALIASES  # noqa: E402

COLS = {"level": 2, "status": 6, "scope": 10, "origin": 11, "disruptive": 12, "health": 13}
# header (row 2 of the Testcases sheet) each column must start with - full_rack_testplan.xlsx layout
HEADERS = {1: "tid", 2: "level", 6: "execution status", 10: "scope", 11: "origin", 12: "disruptive", 13: "health check"}


EMPTY = "(empty)"          # filter value matching a blank cell (e.g. --status "(empty)": never run)
_SPLIT = re.compile(r",(?![^()]*\))")
TAG_SEP = ";"               # a cell may carry several tags: Scope "Switch + Compute; IFCP" is both


def tags(value: str) -> list[str]:
    """The tags of a cell: "Switch + Compute; IFCP" -> ["Switch + Compute", "IFCP"]; a plain value is one tag."""
    return [t.strip() for t in str(value or "").split(TAG_SEP) if t.strip()]


class WorkbookLayoutError(ValueError):
    pass


def check_layout(ws) -> None:
    """Refuse a workbook whose Testcases columns are not where COLS expects them (another plan layout would
    silently filter on the wrong columns)."""
    hdr = next(ws.iter_rows(min_row=2, max_row=2, values_only=True), ())
    bad = [f"column {c}: expected '{want}', found '{(hdr[c - 1] if c <= len(hdr) else None) or ''}'"
           for c, want in HEADERS.items()
           if not str((hdr[c - 1] if c <= len(hdr) else None) or "").strip().lower().startswith(want)]
    if bad:
        raise WorkbookLayoutError("not the full-rack test-plan layout (Testcases row 2): " + "; ".join(bad))


def split_values(v) -> list[str]:
    """Filter values: a comma-separated string (commas inside parentheses belong to the value, e.g. "Health check
    (Switch + Compute, AFM API)") or a list of values. Blank parts dropped."""
    parts = v if isinstance(v, (list, tuple)) else _SPLIT.split(str(v or ""))
    return [str(x).strip() for x in parts if str(x).strip()]


def wanted(filters: dict) -> dict[str, set[str]]:
    """{column: allowed lower-case values}; a column without values is no filter. EMPTY matches a blank cell."""
    out = {}
    for k, v in filters.items():
        vals = {x.lower() for x in split_values(v)}
        if vals:
            out[k] = vals
    return out


def matches(case: dict, want: dict[str, set[str]]) -> bool:
    """A plan row (plan_rows) passes every column filter: case-insensitive exact match of the whole cell or of any
    one of its tags (--scope IFCP matches "Switch + Compute; IFCP", so does --scope "Switch + Compute")."""
    def ok(cell: str, allowed: set[str]) -> bool:
        if not cell:
            return EMPTY.lower() in allowed or EMPTY in allowed
        return cell.lower() in allowed or any(t.lower() in allowed for t in tags(cell))
    return all(ok(case[k], v) for k, v in want.items())


def _code_id(tid, title: str):
    """The row's code ID (shared TIDs: TC-AFM-005 RCCL row = TC-AFM-005A, ...; the TID-less "Not allowed in
    Minirack" row = TC-AFM-035A)."""
    t = title.lower()
    return next((c for c, (a, frag) in ROW_ALIASES.items() if a == (tid or None) and frag.lower() in t), tid)


def plan_rows(wb_path: Path) -> list[dict]:
    """Every test-case row of the Testcases sheet (read only), in sheet order: {id (code ID), tid, title, level,
    status, scope, origin, disruptive, health} - cell values stripped, blank = "". Refuses another layout."""
    wb = openpyxl.load_workbook(wb_path, read_only=True)
    try:
        if "Testcases" not in wb.sheetnames:
            raise WorkbookLayoutError("no 'Testcases' sheet")
        ws = wb["Testcases"]
        check_layout(ws)
        hdr = next(ws.iter_rows(min_row=2, max_row=2, values_only=True), ())
        # every titled column by its header (the UI's "any column" filter): {header: cell text}
        named = [(i, " ".join(str(h).split())) for i, h in enumerate(hdr) if h is not None and str(h).strip()]
        out = []
        for row in ws.iter_rows(min_row=3, values_only=True):
            tid = row[0] if row else None
            title = str((row[2] if len(row) > 2 else None) or "").strip()
            if tid and str(tid).startswith(("TC-", "HC-")):
                code = _code_id(tid, title)
            else:
                code = _code_id(None, title) if not tid and title else None
                if not code:
                    continue
            vals = {k: str((row[c - 1] if c <= len(row) else None) or "").strip() for k, c in COLS.items()}
            cells = {name: str((row[i] if i < len(row) else None) or "").strip() for i, name in named}
            out.append({"id": code, "tid": tid, "title": title, **vals, "cells": cells})
        return out
    finally:
        wb.close()


def suite_of(test_file: str) -> str:
    return "afm" if test_file.startswith("tests/afm/") else "ssh"


def select(wb_path: Path, suite: str = "all", **filters) -> list[str]:
    auto = automated()
    want = wanted(filters)
    out = []
    for case in plan_rows(wb_path):
        if not matches(case, want) or case["id"] not in auto:
            continue
        if suite == "all" or suite == suite_of(auto[case["id"]][0]):
            out.append(case["id"])
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workbook", type=Path)
    for k in COLS:
        ap.add_argument(f"--{k}", default="", help=f"comma-separated values of the {k} column")
    ap.add_argument("--suite", choices=("afm", "ssh", "all"), default="all")
    ap.add_argument("-o", "--output", type=Path, help="write the comma-separated list here (else print it)")
    a = ap.parse_args()
    try:
        ids = select(a.workbook, a.suite, **{k: getattr(a, k) for k in COLS})
    except WorkbookLayoutError as e:
        print(f"{a.workbook}: {e}", file=sys.stderr)
        return 2
    text = ",".join(ids)
    if a.output:
        a.output.write_text(text)
        print(f"{len(ids)} case(s) -> {a.output}")
    else:
        print(text)
    return 0


if __name__ == "__main__":
    sys.exit(main())
