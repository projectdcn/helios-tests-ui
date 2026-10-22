"""tools/update_testplan.find_row: TC-AFM-005A / 048A on their own rows (renamed 21 Oct 2026) and on the shared
BKC-26.12 TIDs of older workbook copies; the committed workbook has no duplicate IDs left."""
from __future__ import annotations

import collections
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))
from update_testplan import find_row  # noqa: E402

ROWS = [("TC-AFM-005", "RCCL Cross Domain"), ("TC-AFM-005", "Failover Policy CRUD"),
        ("TC-AFM-048", "Audit Events"), ("TC-AFM-048", "Techsupport Validation (Sideband Only)")]


def _sheet(rows):
    ws = openpyxl.Workbook().active
    for n, (tid, title) in enumerate(rows, start=3):
        ws.cell(n, 1).value, ws.cell(n, 3).value = tid, title
    return ws


def test_old_shared_tids_resolve_by_title():
    ws = _sheet(ROWS)
    assert [find_row(ws, t) for t in ("TC-AFM-005A", "TC-AFM-005", "TC-AFM-048A", "TC-AFM-048")] == [3, 4, 5, 6]


def test_renamed_rows_resolve_by_their_own_id():
    ws = _sheet([("TC-AFM-005A", "RCCL Cross Domain"), ("TC-AFM-005", "Failover Policy CRUD"),
                 ("TC-AFM-048A", "Audit Events"), ("TC-AFM-048", "Techsupport Validation (Sideband Only)")])
    assert [find_row(ws, t) for t in ("TC-AFM-005A", "TC-AFM-005", "TC-AFM-048A", "TC-AFM-048")] == [3, 4, 5, 6]


def test_committed_workbook_has_unique_ids():
    ws = openpyxl.load_workbook(ROOT / "testplan" / "full_rack_testplan.xlsx", read_only=True)["Testcases"]
    ids = collections.Counter(str(r[0]).strip() for r in ws.iter_rows(min_row=3, values_only=True) if r and r[0])
    assert not [t for t, n in ids.items() if n > 1]
    assert ids["TC-AFM-005A"] == ids["TC-AFM-048A"] == 1
